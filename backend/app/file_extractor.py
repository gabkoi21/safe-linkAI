"""Bounded document/image parsing worker. Never executes uploaded content."""

import base64
import io
import json
import sys
import warnings
import zipfile

from docx import Document
from PIL import Image
from pypdf import PdfReader, filters

from app.analysis_errors import AnalysisError

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARACTERS = 60000
MAX_PAGES = 50
MAX_EXPANDED_BYTES = 20 * 1024 * 1024


def extract(data: bytes, extension: str) -> dict:
    if not data:
        raise AnalysisError("The file is empty.", 422)
    if len(data) > MAX_UPLOAD_BYTES:
        raise AnalysisError("Choose a file no larger than 10 MB.", 413)
    pages = None
    if extension == ".txt":
        encoding = (
            "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        )
        try:
            text = data.decode(encoding)
        except UnicodeError as exc:
            raise AnalysisError(
                "TXT files must contain UTF-8 or UTF-16 text.", 422
            ) from exc
        if any(
            ord(char) < 32 and char not in "\r\n\t\f" for char in text
        ) or data.startswith((b"%PDF-", b"PK\x03\x04", b"MZ")):
            raise AnalysisError("The TXT file contains binary data.", 415)
    elif extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise AnalysisError("The file does not contain a valid PDF.", 415)
        # Bound decompression as well as page count and final extracted text.
        filters.ZLIB_MAX_OUTPUT_LENGTH = 4 * 1024 * 1024
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise AnalysisError("Password-protected PDFs are not supported.", 422)
        pages = len(reader.pages)
        if pages > MAX_PAGES:
            raise AnalysisError("PDF files may contain at most 50 pages.", 422)
        parts = []
        count = 0
        for page in reader.pages:
            contents = page.get_contents()
            if contents and len(contents.get_data()) > 4 * 1024 * 1024:
                raise AnalysisError("The PDF page is too complex to process.", 422)
            part = page.extract_text() or ""
            count += len(part) + 1
            if count > MAX_TEXT_CHARACTERS:
                raise AnalysisError(
                    "The file contains too much text. Choose a shorter document (up to 60,000 characters).",
                    422,
                )
            parts.append(part)
        text = "\n".join(parts)
    elif extension == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
            ):
                raise AnalysisError(
                    "The Word document is too large or complex when expanded.", 422
                )
            if (
                "word/document.xml" not in archive.namelist()
                or "[Content_Types].xml" not in archive.namelist()
            ):
                raise AnalysisError("The file is not a DOCX document.", 415)
        document = Document(io.BytesIO(data))
        # Include tables and headers/footers, not just body paragraphs.
        parts = []

        def collect(container):
            for paragraph in container.paragraphs:
                parts.append(paragraph.text)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        collect(cell)

        collect(document)
        for section in document.sections:
            collect(section.header)
            collect(section.footer)
        text = "\n".join(parts)
    elif extension in (".jpg", ".jpeg", ".png"):
        Image.MAX_IMAGE_PIXELS = 20_000_000
        expected = "PNG" if extension == ".png" else "JPEG"
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as image:
                if image.format != expected:
                    raise AnalysisError(
                        "Image content does not match its filename extension.", 415
                    )
                if image.width * image.height > 20_000_000:
                    raise AnalysisError(
                        "Images may contain at most 20 million pixels.", 422
                    )
                image.verify()
            with Image.open(io.BytesIO(data)) as image:
                image.load()
                normalized = io.BytesIO()
                image.convert("RGB" if expected == "JPEG" else "RGBA").save(
                    normalized, format=expected
                )
                image_bytes = normalized.getvalue()
                if len(image_bytes) > 12 * 1024 * 1024:
                    raise AnalysisError(
                        "The image is too large after decoding. Choose a smaller image.",
                        422,
                    )
        return {
            "image": base64.b64encode(image_bytes).decode(),
            "mime_type": "image/png" if expected == "PNG" else "image/jpeg",
            "text": "",
            "pages": None,
        }
    else:
        raise AnalysisError("Supported files: TXT, PDF, DOCX, JPG, JPEG, and PNG.", 415)
    text = text.strip()
    if not text:
        raise AnalysisError(
            "No readable text was found. For scanned PDFs or image-only documents, upload a JPG or PNG screenshot instead.",
            422,
        )
    if len(text) > MAX_TEXT_CHARACTERS:
        raise AnalysisError(
            "Choose a shorter document with up to 60,000 text characters.", 422
        )
    return {"text": text, "pages": pages, "image": None}


def main() -> None:
    try:
        result = extract(sys.stdin.buffer.read(MAX_UPLOAD_BYTES + 1), sys.argv[1])
        output = {"result": result}
    except AnalysisError as exc:
        output = {"error": str(exc), "status": exc.status_code}
    except Exception:
        output = {
            "error": "The file is damaged, unsupported, or could not be read. Try exporting it again.",
            "status": 422,
        }
    sys.stdout.write(json.dumps(output))


if __name__ == "__main__":
    main()
