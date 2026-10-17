import io
import json
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.analysis import MessageAnalysisResult
from app.analysis_errors import AnalysisError
from app.api.deps import get_db
from app.api.routes.analysis import router
from app.core.config import settings
from app.file_analysis import (
    MAX_UPLOAD_BYTES,
    ExtractionResult,
    analyze_prepared_file,
    prepare_file,
)
from app.models import AnalysisReport, AnalysisStatus, File, Submission
from app.upload_limits import UploadSizeLimitMiddleware

TEXT = "Urgent: Send your password now."
RESULT = MessageAnalysisResult(
    risk_score=95,
    risk_level="high",
    summary="Requests credentials.",
    recommended_action="Do not share your password.",
    findings=[
        {
            "title": "Credential request",
            "description": "Asks for a password.",
            "severity": "high",
            "evidence": "Send your password",
        }
    ],
)


def pdf_bytes(text: str | None = TEXT, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=200)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        contents = DecodedStreamObject()
        contents.set_data(f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(contents)
    if encrypted:
        writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph(TEXT)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table payment request"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def image_bytes(format: str = "PNG") -> bytes:
    image = Image.new("RGB", (500, 100), "white")
    ImageDraw.Draw(image).text((10, 20), TEXT, fill="black")
    output = io.BytesIO()
    image.save(output, format=format)
    return output.getvalue()


class FileEndpointTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        patcher = patch.object(
            settings, "FILE_STORAGE_DIR", Path(self.directory.name) / "uploads"
        )
        patcher.start()
        self.addCleanup(patcher.stop)
        self.engine = create_engine(
            "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
        )
        SQLModel.metadata.create_all(self.engine)
        app = FastAPI()
        app.include_router(router, prefix="/api/v1")
        app.add_middleware(UploadSizeLimitMiddleware)

        def database():
            with Session(self.engine) as session:
                yield session

        app.dependency_overrides[get_db] = database
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.engine.dispose()

    def test_all_supported_extensions_save_private_files_reports_and_findings(self):
        files = [
            ("invoice.txt", TEXT.encode()),
            ("invoice.pdf", pdf_bytes()),
            ("invoice.docx", docx_bytes()),
            ("screen.png", image_bytes()),
            ("screen.jpg", image_bytes("JPEG")),
            ("screen.jpeg", image_bytes("JPEG")),
        ]
        extraction = ExtractionResult(
            method="text_extraction", text=TEXT, limitations=["Content review only."]
        )
        with patch(
            "app.api.routes.analysis.analyze_prepared_file",
            return_value=(RESULT, extraction),
        ):
            for name, data in files:
                with self.subTest(name=name):
                    response = self.client.post(
                        "/api/v1/analysis/file",
                        files={"file": (name, data, "application/octet-stream")},
                    )
                    self.assertEqual(response.status_code, 201, response.text)
                    body = response.json()
                    self.assertEqual(body["file"]["original_name"], name)
                    self.assertEqual(body["file"]["size_bytes"], len(data))
                    self.assertEqual(len(body["file"]["sha256"]), 64)
                    self.assertNotIn("storage_key", body["file"])
                    with Session(self.engine) as session:
                        from uuid import UUID

                        upload = session.get(File, UUID(body["file"]["file_id"]))
                        self.assertEqual(
                            (
                                settings.FILE_STORAGE_DIR / upload.storage_key
                            ).read_bytes(),
                            data,
                        )
                        report = session.get(AnalysisReport, UUID(body["report_id"]))
                        self.assertEqual(report.status, AnalysisStatus.COMPLETED)
                        self.assertEqual(
                            report.analysis_details["extraction"]["text"], TEXT
                        )
                        self.assertEqual(len(report.findings), 1)
                        self.assertIsNone(upload.submission.user_id)

    def test_bad_files_never_call_ai_or_save_uploads(self):
        cases = [
            ("bad.exe", b"MZ"),
            ("empty.txt", b""),
            ("bad.pdf", b"not PDF"),
            ("bad.png", b"not PNG"),
            ("binary.txt", b"\x00\x01"),
        ]
        with patch("app.api.routes.analysis.analyze_prepared_file") as analyze:
            for name, data in cases:
                response = self.client.post(
                    "/api/v1/analysis/file", files={"file": (name, data)}
                )
                self.assertIn(response.status_code, (415, 422), response.text)
            analyze.assert_not_called()
        self.assertFalse(settings.FILE_STORAGE_DIR.exists())
        with Session(self.engine) as session:
            self.assertEqual(session.exec(select(Submission)).all(), [])

    def test_provider_failure_keeps_failed_report_without_demo_result(self):
        with patch(
            "app.api.routes.analysis.analyze_prepared_file",
            side_effect=AnalysisError("AI unavailable.", 502),
        ):
            response = self.client.post(
                "/api/v1/analysis/file", files={"file": ("sample.txt", TEXT.encode())}
            )
        self.assertEqual(response.status_code, 502)
        self.assertNotIn("risk_score", response.json())
        with Session(self.engine) as session:
            report = session.exec(select(AnalysisReport)).one()
            self.assertEqual(report.status, AnalysisStatus.FAILED)
            self.assertIsNone(report.risk_score)
            self.assertEqual(len(session.exec(select(File)).all()), 1)

    def test_database_failure_removes_uncommitted_blob(self):
        with patch.object(
            Session, "commit", side_effect=SQLAlchemyError("private database details")
        ):
            response = self.client.post(
                "/api/v1/analysis/file", files={"file": ("sample.txt", TEXT.encode())}
            )
        self.assertEqual(response.status_code, 503)
        self.assertEqual(list(settings.FILE_STORAGE_DIR.iterdir()), [])
        self.assertNotIn("private database details", response.text)

    def test_oversized_body_rejected_before_parsing(self):
        with patch("app.api.routes.analysis.prepare_file") as prepare:
            response = self.client.post(
                "/api/v1/analysis/file",
                content=b"x",
                headers={
                    "content-length": str(MAX_UPLOAD_BYTES + 100000),
                    "content-type": "multipart/form-data; boundary=test",
                },
            )
            self.assertEqual(response.status_code, 413)
            prepare.assert_not_called()


class FileExtractionTests(unittest.TestCase):
    def test_extracts_txt_pdf_and_docx_tables(self):
        self.assertEqual(prepare_file("message.txt", TEXT.encode()).text, TEXT)
        self.assertEqual(prepare_file("message.txt", TEXT.encode("utf-16")).text, TEXT)
        self.assertIn(TEXT, prepare_file("invoice.pdf", pdf_bytes()).text)
        self.assertIn(
            "Table payment request", prepare_file("invoice.docx", docx_bytes()).text
        )

    def test_pdf_without_text_and_encrypted_pdf_rejected(self):
        for content in (pdf_bytes(None), pdf_bytes(encrypted=True)):
            with self.assertRaises(AnalysisError) as error:
                prepare_file("scan.pdf", content)
            self.assertEqual(error.exception.status_code, 422)

    def test_size_text_and_image_format_limits(self):
        for name, data in (
            ("large.txt", b"x" * (MAX_UPLOAD_BYTES + 1)),
            ("long.txt", b"x" * 60001),
            ("wrong.jpg", image_bytes("PNG")),
        ):
            with self.assertRaises(AnalysisError):
                prepare_file(name, data)

    def test_zip_cannot_impersonate_docx(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("arbitrary.txt", "not a document")
        with self.assertRaises(AnalysisError):
            prepare_file("fake.docx", output.getvalue())

    def test_parser_timeout_becomes_readable_error(self):
        with (
            patch(
                "app.file_analysis.subprocess.run",
                side_effect=subprocess.TimeoutExpired("parser", 15),
            ),
            self.assertRaises(AnalysisError) as error,
        ):
            prepare_file("sample.txt", TEXT.encode())
        self.assertIn("too long", str(error.exception))

    def test_image_is_sent_as_visual_input_and_transcription_is_validated(self):
        prepared = prepare_file("screen.png", image_bytes())
        output = {"transcription": TEXT, "analysis": RESULT.model_dump()}
        with patch(
            "app.file_analysis.call_gemma", return_value=json.dumps(output)
        ) as gemma:
            result, extraction = analyze_prepared_file(prepared)
            self.assertEqual(result.risk_score, 95)
            self.assertEqual(extraction.method, "image_understanding")
            self.assertEqual(
                gemma.call_args.kwargs["contents"][1].inline_data.mime_type, "image/png"
            )
        with (
            patch(
                "app.file_analysis.call_gemma",
                return_value=json.dumps({**output, "transcription": "unrelated text"}),
            ),
            self.assertRaises(AnalysisError),
        ):
            analyze_prepared_file(prepared)
